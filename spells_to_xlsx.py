#!/usr/bin/env python3
"""
spells_to_xlsx.py — OCR the Mage and Illusionist spells PDF and extract
structured spell data into an XLSX workbook using Claude.

Usage:
    python3 spells_to_xlsx.py
    python3 spells_to_xlsx.py --raw-text /tmp/spells_raw.txt
"""

from __future__ import annotations
import argparse
import json
import os
import re
import subprocess
import tempfile
import time
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

import anthropic
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

PDF_PATH  = "/home/user/Test-Game/Mage and Illusionist spells.pdf"
OUT_PATH  = "/home/user/Test-Game/mage_illusionist_spells.xlsx"
MODEL     = "claude-sonnet-4-6"
PAGES_PER_CHUNK = 5

COLUMNS = [
    "name", "class", "level", "school",
    "components", "range", "casting_time", "duration",
    "area_of_effect", "saving_throw", "description",
]

SYSTEM_PROMPT = """\
You are an expert AD&D 1st Edition rules parser. You will receive raw OCR text \
from scanned pages of the AD&D Player's Handbook covering Magic-User and Illusionist spells. \
The text may be garbled, have OCR errors, or mix two columns together.

Your job: identify every COMPLETE spell entry in the text and return a JSON array of spell objects.

Return ONLY a valid JSON array (no prose, no markdown fences). \
If no complete spells are found, return an empty array [].

Each spell object must have ALL of these fields:
{
  "name":          <string, the spell name, title-cased, e.g. "Magic Missile">,
  "class":         <string, "Magic-User", "Illusionist", or "Magic-User/Illusionist" if listed for both>,
  "level":         <integer, spell level 1-9>,
  "school":        <string, school of magic in parentheses after the name, e.g. "Alteration", "Evocation", "Illusion/Phantasm">,
  "components":    <string, e.g. "V, S" or "V, S, M">,
  "range":         <string, e.g. '6"' or "Touch" or "0">,
  "casting_time":  <string, e.g. "1 segment" or "1 round">,
  "duration":      <string, e.g. "1 round/level" or "Permanent" or "Special">,
  "area_of_effect":<string, e.g. "One creature" or '1" radius'>,
  "saving_throw":  <string, e.g. "None" or "Neg." or "½">,
  "description":   <string, the full Explanation/Description text, cleaned of OCR artifacts>
}

Rules:
- Only include spells that appear COMPLETE in this chunk (name + stats + description all present).
- If a spell is cut off (description starts but clearly continues beyond this chunk), omit it — \
  it will be captured in the next chunk.
- The section heading "MAGIC-USER SPELLS" or "ILLUSIONIST SPELLS" tells you the class for \
  all spells that follow until the next section heading.
- Level subheadings like "First Level Spells:" or "MAGIC-USER SPELLS (3RD LEVEL)" indicate \
  the level for the spells that follow.
- Ignore Druid spells, Cleric spells, or any non-Magic-User/Illusionist content.
- Fix obvious OCR errors in spell names and descriptions (e.g. "deoth" → "death").
- Do not invent or add information not present in the text.
"""


def ocr_pdf(pdf_path: str) -> str:
    """OCR the PDF using Tesseract, return full text with page markers."""
    print("Converting PDF to images …")
    with tempfile.TemporaryDirectory() as tmpdir:
        base = os.path.join(tmpdir, "page")
        subprocess.run(
            ["pdftoppm", "-r", "300", "-png", pdf_path, base],
            check=True, capture_output=True,
        )
        pages = sorted(Path(tmpdir).glob("page-*.png"))
        print(f"OCR-ing {len(pages)} pages …")
        parts = []
        for i, png in enumerate(pages, 1):
            text = subprocess.run(
                ["tesseract", str(png), "stdout", "--psm", "6"],
                capture_output=True, text=True,
            ).stdout
            parts.append(f"<!-- page {i} -->\n{text}")
            print(f"\r  {i}/{len(pages)}", end="", flush=True)
        print()
        return "\n\n".join(parts)


def chunk_text(full_text: str, pages_per_chunk: int) -> list[str]:
    """Split OCR text into chunks of N pages, with 1-page overlap."""
    pages = re.split(r"<!-- page \d+ -->", full_text)
    pages = [p.strip() for p in pages if p.strip()]
    chunks = []
    step = pages_per_chunk
    for i in range(0, len(pages), step):
        chunk_pages = pages[max(0, i - 1): i + step]  # 1-page overlap at start
        chunks.append("\n\n".join(chunk_pages))
    return chunks


def extract_spells_from_chunk(client: anthropic.Anthropic,
                               chunk: str, chunk_idx: int) -> list[dict]:
    """Ask Claude to extract all complete spells from a text chunk."""
    user_msg = (
        f"Extract all complete Magic-User and Illusionist spells from this OCR text.\n\n"
        f"--- OCR TEXT ---\n{chunk}\n--- END ---\n\n"
        "Return a JSON array of spell objects now."
    )
    for attempt in range(5):
        try:
            resp = client.messages.create(
                model=MODEL,
                max_tokens=4000,
                system=[{"type": "text", "text": SYSTEM_PROMPT,
                         "cache_control": {"type": "ephemeral"}}],
                messages=[{"role": "user", "content": user_msg}],
            )
            raw = resp.content[0].text.strip()
            raw = re.sub(r"^```[a-z]*\n?", "", raw)
            raw = re.sub(r"\n?```$", "", raw)
            spells = json.loads(raw)
            if isinstance(spells, list):
                return spells
            return []
        except anthropic.RateLimitError:
            wait = 15 * (2 ** attempt)
            print(f"\n  Rate limit chunk {chunk_idx}, retry in {wait}s…", flush=True)
            time.sleep(wait)
        except json.JSONDecodeError as e:
            print(f"\n  JSON error chunk {chunk_idx}: {e}", flush=True)
            return []
    return []


def dedup_spells(spell_list: list[dict]) -> list[dict]:
    """Remove duplicates by name+class+level, keeping the longest description."""
    seen: dict[str, dict] = {}
    for spell in spell_list:
        key = (spell.get("name", "").lower().strip(),
               spell.get("class", ""),
               str(spell.get("level", "")))
        existing = seen.get(key)
        if not existing or len(spell.get("description", "")) > len(existing.get("description", "")):
            seen[key] = spell
    return sorted(seen.values(), key=lambda s: (s.get("class", ""), s.get("level", 0), s.get("name", "")))


def write_xlsx(spells: list[dict], out_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spells"

    # Header row
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    headers = [c.replace("_", " ").title() for c in COLUMNS]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"

    # Column widths
    widths = {
        "name": 30, "class": 22, "level": 7, "school": 22,
        "components": 12, "range": 12, "casting_time": 14, "duration": 20,
        "area_of_effect": 20, "saving_throw": 14, "description": 80,
    }
    for col, col_name in enumerate(COLUMNS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = widths.get(col_name, 15)

    # Alternating row colours
    fill_a = PatternFill("solid", fgColor="DDEEFF")
    fill_b = PatternFill(fill_type=None)

    for row_idx, spell in enumerate(spells, 2):
        fill = fill_a if row_idx % 2 == 0 else fill_b
        for col, col_name in enumerate(COLUMNS, 1):
            val = spell.get(col_name, "")
            if col_name == "level" and val != "":
                try:
                    val = int(val)
                except (ValueError, TypeError):
                    pass
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=(col_name == "description"))
            cell.fill = fill

        # Set description row height
        ws.row_dimensions[row_idx].height = 60

    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw-text", default=None,
                    help="Path to pre-OCR'd text file (skip OCR step)")
    ap.add_argument("--output", default=OUT_PATH)
    ap.add_argument("--workers", type=int, default=3)
    args = ap.parse_args()

    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise SystemExit("ERROR: ANTHROPIC_API_KEY not set")

    # Step 1: get OCR text
    if args.raw_text and os.path.exists(args.raw_text):
        print(f"Using pre-OCR'd text from {args.raw_text}")
        full_text = open(args.raw_text).read()
    else:
        full_text = ocr_pdf(PDF_PATH)
        with open("/tmp/spells_raw.txt", "w") as f:
            f.write(full_text)
        print("Saved raw OCR to /tmp/spells_raw.txt")

    # Step 2: chunk
    chunks = chunk_text(full_text, PAGES_PER_CHUNK)
    print(f"\nProcessing {len(chunks)} chunks with {args.workers} workers …\n")

    client = anthropic.Anthropic(api_key=api_key)
    all_spells: list[dict] = []

    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = {pool.submit(extract_spells_from_chunk, client, chunk, i): i
                   for i, chunk in enumerate(chunks, 1)}
        for fut in as_completed(futures):
            chunk_idx = futures[fut]
            spells = fut.result()
            all_spells.extend(spells)
            print(f"  Chunk {chunk_idx}: {len(spells)} spells found  "
                  f"(total so far: {len(all_spells)})", flush=True)

    # Step 3: dedup and sort
    spells = dedup_spells(all_spells)
    print(f"\nTotal unique spells after dedup: {len(spells)}")

    # Step 4: write XLSX
    print(f"Writing {args.output} …")
    write_xlsx(spells, args.output)
    print(f"Done → {args.output}")

    # Summary
    mu = sum(1 for s in spells if "Magic-User" in s.get("class", ""))
    il = sum(1 for s in spells if "Illusionist" in s.get("class", ""))
    print(f"  Magic-User spells: {mu}")
    print(f"  Illusionist spells: {il}")


if __name__ == "__main__":
    main()
