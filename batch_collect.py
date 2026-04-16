#!/usr/bin/env python3
"""
batch_collect.py — Poll for Batch API results and write them into the XLSX.

Reads batch_id.txt, waits until the batch is complete, then writes
all extracted effects back into encyclopedia_magica_vol1.xlsx.

Usage:
    python3 batch_collect.py
    python3 batch_collect.py --output encyclopedia_magica_vol1.xlsx
"""

from __future__ import annotations
import argparse
import json
import os
import re
import time

import anthropic
import openpyxl
from openpyxl.styles import PatternFill, Alignment

from extract_effects import (
    SYSTEM_PROMPT, EFFECT_COLS, XLSX_PATH, OUT_PATH,
    col_index, _empty_effects,
)

BATCH_ID_FILE = "/home/user/Test-Game/batch_id.txt"
POLL_INTERVAL = 60  # seconds between status checks


def parse_effects(raw: str) -> dict:
    raw = raw.strip()
    raw = re.sub(r"^```[a-z]*\n?", "", raw)
    raw = re.sub(r"\n?```$", "", raw)
    try:
        return json.loads(raw)
    except json.JSONDecodeError as e:
        print(f"  JSON parse error: {e}")
        return _empty_effects(needs_review=True)


def coerce_value(col_name: str, value) -> object:
    str_cols = ("granted_spells", "granted_abilities", "conditions", "notes")
    if value is None or value == "null" or value == "None":
        return -1 if col_name == "ac_set" else ("" if col_name in str_cols else 0)
    if col_name == "ac_set" and not isinstance(value, int):
        try:
            return int(value)
        except (TypeError, ValueError):
            return -1
    return value


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input",  default=XLSX_PATH)
    ap.add_argument("--output", default=OUT_PATH)
    args = ap.parse_args()

    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise SystemExit("ERROR: ANTHROPIC_API_KEY not set")

    if not os.path.exists(BATCH_ID_FILE):
        raise SystemExit(f"ERROR: {BATCH_ID_FILE} not found — run batch_submit.py first")

    batch_id = open(BATCH_ID_FILE).read().strip()
    print(f"Batch ID: {batch_id}")

    client = anthropic.Anthropic(api_key=api_key)

    # Poll until complete
    while True:
        batch = client.messages.batches.retrieve(batch_id)
        counts = batch.request_counts
        print(f"  Status: {batch.processing_status}  "
              f"processing={counts.processing}  succeeded={counts.succeeded}  "
              f"errored={counts.errored}", flush=True)
        if batch.processing_status == "ended":
            break
        time.sleep(POLL_INTERVAL)

    print(f"\nBatch complete. succeeded={counts.succeeded} errored={counts.errored}")
    print(f"Loading {args.input} …")
    wb = openpyxl.load_workbook(args.input)
    ws = wb.active

    col_map = {h: col_index(ws, h) for h in
               ["name", "needs_review"] + EFFECT_COLS}

    review_fill = PatternFill("solid", fgColor="FFF2CC")
    normal_fill = PatternFill(fill_type=None)

    written = 0
    errors  = 0

    for result in client.messages.batches.results(batch_id):
        custom_id = result.custom_id  # e.g. "row-2"
        try:
            row_num = int(custom_id.split("-")[1])
        except (IndexError, ValueError):
            print(f"  Skipping unrecognised custom_id: {custom_id}")
            continue

        if result.result.type == "succeeded":
            raw = result.result.message.content[0].text
            effects = parse_effects(raw)
        else:
            print(f"  {custom_id}: {result.result.type} — marking needs_review")
            effects = _empty_effects(needs_review=True)
            errors += 1

        # Write into the correct Excel row
        row = list(ws.iter_rows(min_row=row_num, max_row=row_num, values_only=False))[0]
        for col_name in EFFECT_COLS:
            col_idx = col_map.get(col_name)
            if col_idx is None:
                continue
            cell = row[col_idx - 1]
            cell.value = coerce_value(col_name, effects.get(col_name))
            cell.alignment = Alignment(vertical="top", wrap_text=False)

        nr = effects.get("needs_review", False)
        fill = review_fill if nr else normal_fill
        for cell in row:
            cell.fill = fill

        written += 1
        if written % 100 == 0:
            print(f"  Written {written} rows …", flush=True)

    print(f"\nSaving {args.output} …")
    wb.save(args.output)

    n_review = sum(
        1 for row in ws.iter_rows(min_row=2, values_only=True)
        if row[col_map["needs_review"] - 1]
    )
    print(f"Done → {args.output}")
    print(f"Rows written: {written}  Errors: {errors}  needs_review: {n_review}")


if __name__ == "__main__":
    main()
