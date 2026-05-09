#!/usr/bin/env python3
"""
combine_spells.py

Reads multiple spell JSON files, deduplicates, sorts, and writes a formatted XLSX.
"""

import json
import os
import sys

try:
    import openpyxl
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl is not installed. Run: pip3 install openpyxl")
    sys.exit(1)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

INPUT_FILES = [
    "/tmp/spells_mu_1_2.json",
    "/tmp/spells_mu_3_4.json",
    "/tmp/spells_mu_5.json",
    "/tmp/spells_mu_6.json",
    "/tmp/spells_mu_7.json",
    "/tmp/spells_mu_8.json",
    "/tmp/spells_mu_9.json",
    "/tmp/spells_ill_1_3.json",
    "/tmp/spells_ill_4_5.json",
    "/tmp/spells_ill_6_7.json",
    "/tmp/spells_illusionist.json",
]

OUTPUT_FILE = "/home/user/Test-Game/mage_illusionist_spells.xlsx"

COLUMNS = [
    "name",
    "class",
    "level",
    "school",
    "components",
    "range",
    "casting_time",
    "duration",
    "area_of_effect",
    "saving_throw",
    "description",
]

COLUMN_WIDTHS = {
    "name": 30,
    "class": 22,
    "level": 7,
    "school": 22,
    "components": 12,
    "range": 12,
    "casting_time": 14,
    "duration": 20,
    "area_of_effect": 20,
    "saving_throw": 14,
    "description": 80,
}

# Formatting constants
HEADER_FILL_COLOR = "1F4E79"   # Dark blue
HEADER_FONT_COLOR = "FFFFFF"   # White
EVEN_ROW_FILL_COLOR = "DDEEFF" # Light blue
DATA_ROW_HEIGHT = 60

# Class sort order
CLASS_ORDER = {"Magic-User": 0, "Illusionist": 1}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def load_spells_from_file(filepath):
    """Load spells from a JSON file. Returns a list of dicts."""
    try:
        with open(filepath, "r", encoding="utf-8") as fh:
            data = json.load(fh)
    except FileNotFoundError:
        print(f"  [skip] {filepath} — file not found")
        return []
    except json.JSONDecodeError as exc:
        print(f"  [skip] {filepath} — JSON parse error: {exc}")
        return []

    # Accept either a bare list or {"spells": [...]}
    if isinstance(data, list):
        spells = data
    elif isinstance(data, dict):
        # Try common wrapper keys
        for key in ("spells", "data", "results"):
            if key in data and isinstance(data[key], list):
                spells = data[key]
                break
        else:
            # Treat the dict itself as a single spell if it has a "name" key
            if "name" in data:
                spells = [data]
            else:
                print(f"  [skip] {filepath} — unrecognised JSON structure")
                return []
    else:
        print(f"  [skip] {filepath} — unexpected JSON type: {type(data)}")
        return []

    print(f"  [ok]   {filepath} — {len(spells)} spell(s) loaded")
    return spells


def dedup_key(spell):
    """Return the deduplication key tuple."""
    name = str(spell.get("name", "")).strip().lower()
    cls  = str(spell.get("class", "")).strip()
    lvl  = spell.get("level", 0)
    return (name, cls, lvl)


def deduplicate(spells):
    """
    Deduplicate by (name.lower(), class, level).
    Among duplicates, keep the entry with the longest description.
    """
    best = {}
    for spell in spells:
        key = dedup_key(spell)
        desc_len = len(str(spell.get("description", "")))
        if key not in best or desc_len > len(str(best[key].get("description", ""))):
            best[key] = spell
    return list(best.values())


def sort_key(spell):
    """Sort: class (Magic-User=0, Illusionist=1, others=2), level, name."""
    cls  = str(spell.get("class", "")).strip()
    lvl  = spell.get("level", 0)
    name = str(spell.get("name", "")).strip().lower()
    return (CLASS_ORDER.get(cls, 2), lvl, name)


def cell_value(spell, col):
    """Return the value for a given column, coercing to an appropriate type."""
    val = spell.get(col, "")
    if val is None:
        return ""
    if col == "level":
        try:
            return int(val)
        except (ValueError, TypeError):
            return val
    return str(val)


# ---------------------------------------------------------------------------
# XLSX writing
# ---------------------------------------------------------------------------

def write_xlsx(spells, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spells"

    # --- Header row ---
    header_fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL_COLOR)
    header_font = Font(bold=True, color=HEADER_FONT_COLOR)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name.replace("_", " ").title())
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # Freeze pane at A2 (freeze the header row)
    ws.freeze_panes = "A2"

    # --- Column widths ---
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = COLUMN_WIDTHS.get(col_name, 15)

    # --- Data rows ---
    even_fill = PatternFill(fill_type="solid", fgColor=EVEN_ROW_FILL_COLOR)
    no_fill   = PatternFill(fill_type=None)  # No fill for odd rows

    desc_col_idx = COLUMNS.index("description") + 1  # 1-based

    for row_idx, spell in enumerate(spells, start=2):
        is_even = (row_idx % 2 == 0)
        row_fill = even_fill if is_even else no_fill

        for col_idx, col_name in enumerate(COLUMNS, start=1):
            val = cell_value(spell, col_name)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)

            # Fill
            cell.fill = row_fill

            # Alignment
            if col_idx == desc_col_idx:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=False)

        # Row height
        ws.row_dimensions[row_idx].height = DATA_ROW_HEIGHT

    wb.save(output_path)
    print(f"\nXLSX written to: {output_path}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    print("=== combine_spells.py ===\n")

    # 1. Load all spells
    all_spells = []
    for filepath in INPUT_FILES:
        loaded = load_spells_from_file(filepath)
        all_spells.extend(loaded)

    print(f"\nTotal spells loaded (before dedup): {len(all_spells)}")

    if not all_spells:
        print("\nNo spells were loaded — nothing to write.")
        sys.exit(0)

    # 2. Deduplicate
    spells = deduplicate(all_spells)
    print(f"Total spells after deduplication:   {len(spells)}")

    # 3. Sort
    spells.sort(key=sort_key)

    # 4. Write XLSX
    write_xlsx(spells, OUTPUT_FILE)

    # 5. Summary
    mu_count   = sum(1 for s in spells if str(s.get("class", "")).strip() == "Magic-User")
    ill_count  = sum(1 for s in spells if str(s.get("class", "")).strip() == "Illusionist")
    other_count = len(spells) - mu_count - ill_count

    print("\n--- Summary ---")
    print(f"  Total spells:      {len(spells)}")
    print(f"  Magic-User spells: {mu_count}")
    print(f"  Illusionist spells:{ill_count}")
    if other_count:
        print(f"  Other/unknown:     {other_count}")


if __name__ == "__main__":
    main()
